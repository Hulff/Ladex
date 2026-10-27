import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tabconv.core.errors import ExportError, OperationCancelledError
from tabconv.core.export import write_dataset
from tabconv.core.models import Dataset, ExportOptions


@pytest.fixture
def dataset() -> Dataset:
    return Dataset(("name", "value"), (("=literal", "two, words"), ("second", "")))


@pytest.mark.parametrize("suffix,delimiter", [("csv", ","), ("tsv", "\t")])
def test_export_and_reimport_text(tmp_path: Path, dataset: Dataset, suffix: str, delimiter: str) -> None:
    path = tmp_path / f"result.{suffix}"
    write_dataset(dataset, path, ExportOptions(delimiter=delimiter, line_ending="\r\n"))
    assert path.read_bytes().count(b"\r\n") == 3
    with path.open(encoding="utf-8", newline="") as stream:
        assert list(csv.reader(stream, delimiter=delimiter)) == [
            list(dataset.columns),
            list(dataset.rows[0]),
            list(dataset.rows[1]),
        ]


def test_xlsx_cells_are_written_as_strings(tmp_path: Path, dataset: Dataset) -> None:
    path = tmp_path / "result.xlsx"
    write_dataset(dataset, path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Data"]
    assert list(sheet.values) == [("name", "value"), ("=literal", "two, words"), ("second", None)]
    assert sheet.cell(row=2, column=1).data_type == "s"
    workbook.close()


def test_existing_file_is_not_overwritten_without_confirmation(
    tmp_path: Path, dataset: Dataset
) -> None:
    path = tmp_path / "result.csv"
    path.write_text("keep me")
    with pytest.raises(ExportError):
        write_dataset(dataset, path)
    assert path.read_text() == "keep me"


def test_explicit_overwrite_replaces_atomically(tmp_path: Path, dataset: Dataset) -> None:
    path = tmp_path / "result.csv"
    path.write_text("old")
    write_dataset(dataset, path, overwrite=True)
    assert path.read_text().startswith("name,value")
    assert not list(tmp_path.glob(".result.*"))


def test_cancelled_export_leaves_no_partial_file(tmp_path: Path, dataset: Dataset) -> None:
    path = tmp_path / "result.csv"
    with pytest.raises(OperationCancelledError):
        write_dataset(dataset, path, cancelled=lambda: True)
    assert not path.exists()
    assert not list(tmp_path.glob(".result.*"))


def test_export_does_not_create_missing_parent(tmp_path: Path, dataset: Dataset) -> None:
    with pytest.raises(ExportError):
        write_dataset(dataset, tmp_path / "missing" / "result.csv")
